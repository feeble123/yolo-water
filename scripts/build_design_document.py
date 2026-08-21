"""Build the official design DOCX from the retained template and Markdown source."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

TEMPLATE_SHA256 = "8e8148db47a1fd7050c11027105f85299c72596ffaefd4c9a700824afc9877d0"
TABLE_WIDTH_DXA = 8180


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _clear_content(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _font(run, *, size: float = 10.5, bold: bool = False) -> None:
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), "Arial")


def _format_paragraph(paragraph, *, before: float = 3, after: float = 6) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.35
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _metadata(paragraph, label: str, value: str) -> None:
    _clear_content(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _format_paragraph(paragraph, before=6, after=6)
    _font(paragraph.add_run(label), size=11, bold=True)
    _font(paragraph.add_run(f"：{value}"), size=11)


def _remove_template_body(document: Document) -> None:
    body = document._element.body
    keep = {paragraph._p for paragraph in document.paragraphs[:4]}
    for child in list(body):
        if child.tag == qn("w:sectPr") or child in keep:
            continue
        body.remove(child)


def _chapter(document: Document, text: str, *, page_break: bool) -> None:
    paragraph = document.add_paragraph()
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(16)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.2
    _font(paragraph.add_run(text), size=16, bold=True)


def _subheading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.2
    _font(paragraph.add_run(text), size=12, bold=True)


def _body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    _format_paragraph(paragraph)
    _font(paragraph.add_run(text))


def _list_item(document: Document, text: str, *, numbered: bool) -> None:
    # The official template contains no numbering part or list styles. Render source
    # list items as grouped, indented prose instead of fabricating Unicode bullets or
    # manual numbering. The source order still preserves procedural sequencing.
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(12 if numbered else 9)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    label, separator, remainder = text.partition("：")
    if separator and len(label) <= 16:
        _font(paragraph.add_run(f"{label}："), bold=True)
        _font(paragraph.add_run(remainder))
    else:
        _font(paragraph.add_run(text))


def _cell_margins(cell, top: int = 100, start: int = 100, bottom: int = 100, end: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_width(element, width: int) -> None:
    element.set(qn("w:w"), str(width))
    element.set(qn("w:type"), "dxa")


def _table(document: Document, rows: list[list[str]]) -> None:
    column_count = len(rows[0])
    if column_count == 2:
        weights = [0.36, 0.64]
    elif column_count == 3:
        weights = [0.28, 0.28, 0.44]
    else:
        weights = [1 / column_count] * column_count
    widths = [round(TABLE_WIDTH_DXA * weight) for weight in weights]
    widths[-1] += TABLE_WIDTH_DXA - sum(widths)

    table = document.add_table(rows=len(rows), cols=column_count)
    table.autofit = False
    table_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "B7C9D3")
        borders.append(border)
    table_pr.append(borders)
    table_width = table_pr.first_child_found_in("w:tblW")
    _set_width(table_width, TABLE_WIDTH_DXA)
    table_indent = OxmlElement("w:tblInd")
    _set_width(table_indent, 100)
    table_pr.append(table_indent)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _cell_margins(cell)
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            _set_width(tc_width, widths[column_index])
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if column_index < 2 else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.15
            _font(paragraph.add_run(value), size=9.5, bold=row_index == 0)
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "D9EAF0")
                cell._tc.get_or_add_tcPr().append(shading)
        if row_index == 0:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            table.rows[0]._tr.get_or_add_trPr().append(header)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _parse_table(lines: list[str]) -> list[list[str]]:
    return [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]


def _is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _render_markdown(document: Document, source: str) -> None:
    lines = source.splitlines()
    index = 0
    chapter_count = 0
    paragraph_buffer: list[str] = []

    def flush() -> None:
        if paragraph_buffer:
            _body(document, " ".join(paragraph_buffer))
            paragraph_buffer.clear()

    while index < len(lines):
        line = lines[index].strip()
        if not line:
            flush()
            index += 1
            continue
        if line.startswith("## "):
            flush()
            chapter_count += 1
            _chapter(document, line[3:].strip(), page_break=chapter_count > 1)
            index += 1
            continue
        if line.startswith("### "):
            flush()
            _subheading(document, line[4:].strip())
            index += 1
            continue
        if line.startswith(("# ", "作品名称：")):
            flush()
            index += 1
            continue
        if (
            line.startswith("| ")
            and index + 1 < len(lines)
            and _is_table_separator(lines[index + 1])
        ):
            flush()
            table_lines = [line]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            _table(document, _parse_table(table_lines))
            continue
        if line.startswith("- "):
            flush()
            _list_item(document, line[2:].strip(), numbered=False)
            index += 1
            continue
        match = re.match(r"^\d+\.\s+(.+)$", line)
        if match:
            flush()
            _list_item(document, match.group(1), numbered=True)
            index += 1
            continue
        paragraph_buffer.append(line)
        index += 1
    flush()


def build(
    template: Path,
    source: Path,
    output: Path,
    team_name: str,
    completed: str,
    *,
    force: bool,
) -> dict[str, object]:
    if _hash(template) != TEMPLATE_SHA256:
        raise ValueError("官方模板哈希已变化，请重新执行模板蒸馏")
    if output.resolve() == template.resolve():
        raise ValueError("输出路径不能覆盖官方原始模板")
    if output.exists() and not force:
        raise FileExistsError(f"拒绝覆盖已有方案文档：{output}")
    markdown = source.read_text(encoding="utf-8")
    if "## 一、需求与痛点分析" not in markdown or "## 五、效果评估" not in markdown:
        raise ValueError("方案源文档缺少官方五章结构")

    document = Document(template)
    _metadata(document.paragraphs[1], "参赛队伍", team_name)
    _metadata(
        document.paragraphs[2],
        "作品名称",
        "水域智巡——基于YOLO与Qwen的水域综合异常识别智能体",
    )
    _metadata(document.paragraphs[3], "完成日期", completed)
    _remove_template_body(document)
    _render_markdown(document, markdown)
    document.core_properties.title = "水域智巡智能体设计方案"
    document.core_properties.author = "参赛队伍"
    document.core_properties.subject = "水域综合异常识别行业智能体"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)

    if _hash(template) != TEMPLATE_SHA256:
        raise RuntimeError("原始模板在构建过程中发生变化")
    return {
        "output": str(output.resolve()),
        "sha256": _hash(output),
        "team_name": team_name,
        "source": str(source.resolve()),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--team-name", default="【提交前填写报名队伍名】")
    parser.add_argument("--date", default=datetime.now(UTC).date().isoformat())
    parser.add_argument("--force", action="store_true", help="显式覆盖既有生成版方案文档")
    args = parser.parse_args()
    print(
        json.dumps(
            build(
                args.template,
                args.source,
                args.output,
                args.team_name,
                args.date,
                force=args.force,
            ),
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
