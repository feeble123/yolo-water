"""Read-only structural inspection for the official DOCX design template."""

from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document


def _run_info(run) -> dict[str, object]:
    color = run.font.color.rgb
    return {
        "text": run.text,
        "font": run.font.name,
        "size_pt": run.font.size.pt if run.font.size else None,
        "bold": run.bold,
        "italic": run.italic,
        "color": str(color) if color else None,
    }


def inspect(path: Path) -> dict[str, object]:
    document = Document(path)
    report: dict[str, object] = {
        "reference": str(path.resolve()),
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "paragraphs": [],
        "tables": [],
        "sections": [],
    }

    report["paragraphs"] = [
        {
            "index": index,
            "text": paragraph.text,
            "style": paragraph.style.name,
            "alignment": str(paragraph.alignment),
            "space_before_pt": (
                paragraph.paragraph_format.space_before.pt
                if paragraph.paragraph_format.space_before
                else None
            ),
            "space_after_pt": (
                paragraph.paragraph_format.space_after.pt
                if paragraph.paragraph_format.space_after
                else None
            ),
            "line_spacing": str(paragraph.paragraph_format.line_spacing),
            "runs": [_run_info(run) for run in paragraph.runs],
        }
        for index, paragraph in enumerate(document.paragraphs)
    ]
    report["tables"] = [
        {
            "index": index,
            "style": table.style.name if table.style else None,
            "rows": [[cell.text for cell in row.cells] for row in table.rows],
        }
        for index, table in enumerate(document.tables)
    ]
    report["sections"] = [
        {
            "width_in": section.page_width.inches,
            "height_in": section.page_height.inches,
            "left_in": section.left_margin.inches,
            "right_in": section.right_margin.inches,
            "top_in": section.top_margin.inches,
            "bottom_in": section.bottom_margin.inches,
            "header_in": section.header_distance.inches,
            "footer_in": section.footer_distance.inches,
        }
        for section in document.sections
    ]
    with zipfile.ZipFile(path) as archive:
        report["package_parts"] = [
            {"name": info.filename, "size": info.file_size}
            for info in archive.infolist()
        ]
    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()
    print(json.dumps(inspect(args.docx), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
